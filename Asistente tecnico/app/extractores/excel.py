"""Extracción de hojas y valores desde libros Excel."""

from __future__ import annotations

from pathlib import Path


def extraer_excel(ruta: Path) -> tuple[str, dict[str, str]]:
    """Devuelve las celdas con valor de cada hoja, conservando su posición lógica."""
    if ruta.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise RuntimeError("El formato .xls antiguo no está soportado por el extractor actual.")
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Falta instalar 'openpyxl' para extraer Excel.") from error

    libro = load_workbook(filename=ruta, read_only=True, data_only=True, keep_links=False)
    cantidad_hojas = len(libro.sheetnames)
    bloques: list[str] = []
    try:
        for hoja in libro.worksheets:
            filas: list[str] = []
            for numero, valores in enumerate(hoja.iter_rows(values_only=True), start=1):
                valores_limpios = list(valores)
                while valores_limpios and valores_limpios[-1] is None:
                    valores_limpios.pop()
                if not valores_limpios or all(valor is None for valor in valores_limpios):
                    continue
                celdas = [f"{hoja.cell(numero, columna).coordinate}={valor}" for columna, valor in enumerate(valores_limpios, start=1)]
                filas.append(" | ".join(celdas))
            if filas:
                bloques.append(f"[Hoja: {hoja.title}]\n" + "\n".join(filas))
    finally:
        libro.close()

    contenido = "\n\n".join(bloques) or "[Libro Excel sin valores extraíbles.]"
    return contenido, {"tipo": "excel", "hojas": str(cantidad_hojas)}
